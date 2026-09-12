# -*- coding: utf-8 -*-
"""
Скрипт автоматического аудита и валидации презентации дипломного проекта.
Проверяет:
1. Количество слайдов (18).
2. Формат 16:9 (13.333 x 7.5 дюймов).
3. Наличие заметок докладчика на каждом слайде.
4. Отсутствие незаполненных плейсхолдеров (Lorem, TODO, [Вставьте]).
5. Наличие обязательных инженерных метрик (62 теста, 945.6 RPS, 115-ФЗ, ст. 395 ГК РФ).
"""

import os
import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from pptx import Presentation

REQUIRED_SLIDES = 18
REQUIRED_METRICS = [
    ("62", "62 автотеста pytest"),
    ("945", "945.6 RPS нагрузочный тест Locust"),
    ("115-ФЗ", "Федеральный закон № 115-ФЗ"),
    ("395", "Статья 395 ГК РФ"),
    ("Whitelist", "Механизм Action Whitelist"),
    ("RAG", "Семантический поиск RAG"),
    ("FastAPI", "Интеграционный слой FastAPI")
]

FORBIDDEN_PATTERNS = ["lorem ipsum", "[вставьте", "[вставить", "todo", "xxx"]

def validate_presentation(file_path):
    if not os.path.exists(file_path):
        print(f"[ОШИБКА] Файл презентации не найден: {file_path}")
        return False

    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    errors = []
    warnings = []

    print(f"=== АУДИТ ПРЕЗЕНТАЦИИ: {file_path} ===")

    # 1. Проверка количества слайдов
    if total_slides != REQUIRED_SLIDES:
        errors.append(f"Количество слайдов: {total_slides} (требуется ровно {REQUIRED_SLIDES})")
    else:
        print(f"[OK] Количество слайдов: {total_slides} из {REQUIRED_SLIDES}")

    # 2. Проверка пропорций 16:9
    w = round(prs.slide_width.inches, 3)
    h = round(prs.slide_height.inches, 3)
    ratio = w / h if h > 0 else 0
    if abs(ratio - 16 / 9) > 0.05:
        errors.append(f"Неверное соотношение сторон: {w}x{h} (соотношение {ratio:.2f}, требуется 16:9 / 1.78)")
    else:
        print(f"[OK] Соотношение сторон: 16:9 ({w}'' x {h}'')")

    # 3. Проверка каждого слайда
    all_text_lower = ""
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        slide_texts.append(t)
                        all_text_lower += " " + t.lower()
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            t = p.text.strip()
                            if t:
                                slide_texts.append(t)
                                all_text_lower += " " + t.lower()

        # Проверка заметок докладчика
        notes_text = ""
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip()
            all_text_lower += " " + notes_text.lower()

        if not notes_text:
            warnings.append(f"Слайд {slide_num}: отсутствуют заметки для докладчика (Speaker Notes)")

        # Проверка на запрещенные плейсхолдеры
        joined_text = " ".join(slide_texts).lower()
        for pat in FORBIDDEN_PATTERNS:
            if pat in joined_text:
                errors.append(f"Слайд {slide_num}: обнаружен запрещенный плейсхолдер '{pat}'")

    # 4. Проверка ключевых метрик
    for token, desc in REQUIRED_METRICS:
        if token.lower() in all_text_lower:
            print(f"[OK] Метрика найдена: {desc}")
        else:
            warnings.append(f"Не найдено упоминание обязательной метрики: {desc}")

    print("\n--- ИТОГИ ВАЛИДАЦИИ ---")
    if errors:
        print(f"[FAIL] Выявлено критических ошибок: {len(errors)}")
        for e in errors:
            print(f"  ❌ {e}")
    else:
        print("[SUCCESS] Критических ошибок не обнаружено!")

    if warnings:
        print(f"[ПРЕДУПРЕЖДЕНИЕ] Замечаний: {len(warnings)}")
        for w in warnings:
            print(f"  ⚠️ {w}")
    else:
        print("[SUCCESS] Замечаний нет!")

    return len(errors) == 0


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join("release", "Интеллектуальный_ассистент_для_1С_УНФ_КОНКУРС.pptx")
    success = validate_presentation(target)
    sys.exit(0 if success else 1)
