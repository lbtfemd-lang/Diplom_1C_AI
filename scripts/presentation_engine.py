# -*- coding: utf-8 -*-
"""
Графический движок presentation_engine.py
Реализует дизайн-систему Modern Dark Enterprise для конкурса фирмы «1С».
"""

from typing import List, Dict, Optional, Any
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN


# ==============================================================================
# ЦВЕТОВАЯ ПАЛИТРА MODERN DARK ENTERPRISE
# ==============================================================================
COLOR_CANVAS_BG = RGBColor(0x0B, 0x0F, 0x19)       # #0B0F19 Глубокий темный сланец
COLOR_CARD_BG = RGBColor(0x1E, 0x29, 0x3B)         # #1E293B Темно-синяя карточка
COLOR_CARD_BG_ALT = RGBColor(0x11, 0x18, 0x27)     # #111827 Углубленная подложка
COLOR_CARD_BORDER = RGBColor(0x33, 0x41, 0x55)     # #334155 Slate 700
COLOR_CARD_BORDER_LIGHT = RGBColor(0x47, 0x55, 0x69) # #475569 Slate 600

COLOR_BRAND_1C = RGBColor(0xFF, 0xB3, 0x00)        # #FFB300 Янтарно-золотой акцент «1С»
COLOR_BRAND_YELLOW = RGBColor(0xFF, 0xE6, 0x00)    # #FFE600 Яркий предупреждающий
COLOR_TECH_CYAN = RGBColor(0x06, 0xB6, 0xD4)       # #06B6D4 Неоновый циан
COLOR_TECH_INDIGO = RGBColor(0x63, 0x66, 0xF1)     # #6366F1 Индиго для архитектурных связей
COLOR_SUCCESS_GREEN = RGBColor(0x10, 0xB9, 0x81)   # #10B981 Изумрудный статус
COLOR_DANGER_RED = RGBColor(0xEF, 0x44, 0x44)      # #EF4444 Рубиновый стоп-фактор

COLOR_TEXT_PRIMARY = RGBColor(0xFF, 0xFF, 0xFF)    # #FFFFFF Белый 100% контраст
COLOR_TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)      # #94A3B8 Серебристо-пепельный
COLOR_TEXT_SUBTLE = RGBColor(0x64, 0x74, 0x8B)     # #64748B Slate 500
COLOR_DIVIDER = RGBColor(0x33, 0x41, 0x55)         # #334155 Линия-разделитель

# Шрифты дизайн-системы
FONT_TITLE = "Montserrat"
FONT_BODY = "Inter"
FONT_MONO = "JetBrains Mono"


class PresentationEngine:
    """Декларативный движок верстки презентации 16:9 в стиле Modern Dark Enterprise."""

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]

    def create_slide(
        self,
        title: str,
        subtitle: Optional[str] = None,
        category_badge: Optional[str] = None,
        badge_color: RGBColor = COLOR_TECH_CYAN,
        slide_num: Optional[int] = None
    ) -> Any:
        """Создает базовый слайд с темным фоном, шапкой и линией-разделителем."""
        slide = self.prs.slides.add_slide(self.blank_layout)

        # 1. Фоновый холст
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_CANVAS_BG
        bg.line.fill.background()

        # 2. Категорийный бейдж в верхнем левом углу
        top_offset = Inches(0.4)
        if category_badge:
            self.add_badge(
                slide,
                left=Inches(0.8),
                top=top_offset,
                text=category_badge.upper(),
                bg_color=RGBColor(0x1E, 0x29, 0x3B),
                text_color=badge_color,
                font_size=Pt(9),
                height=Inches(0.28)
            )
            top_offset += Inches(0.36)

        # 3. Заголовок слайда
        title_box = slide.shapes.add_textbox(
            Inches(0.8), top_offset, Inches(11.733), Inches(0.55)
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_bottom = Inches(0)
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY

        # 4. Подзаголовок (если задан)
        if subtitle:
            sub_p = tf.add_paragraph()
            sub_p.text = subtitle
            sub_p.font.name = FONT_BODY
            sub_p.font.size = Pt(12)
            sub_p.font.color.rgb = COLOR_TEXT_MUTED
            sub_p.space_before = Pt(4)

        # 5. Разделитель заголовка
        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.02)
        )
        divider.fill.solid()
        divider.fill.fore_color.rgb = COLOR_DIVIDER
        divider.line.fill.background()

        # 6. Номер слайда в правом нижнем углу
        if slide_num is not None:
            num_box = slide.shapes.add_textbox(
                Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.3)
            )
            ntf = num_box.text_frame
            ntf.margin_left = 0
            ntf.margin_right = 0
            np = ntf.paragraphs[0]
            np.text = f"{slide_num:02d} / 18"
            np.alignment = PP_ALIGN.RIGHT
            np.font.name = FONT_MONO
            np.font.size = Pt(10)
            np.font.color.rgb = COLOR_TEXT_SUBTLE

        return slide

    def add_badge(
        self,
        slide: Any,
        left: Inches,
        top: Inches,
        text: str,
        bg_color: RGBColor = COLOR_CARD_BG,
        text_color: RGBColor = COLOR_BRAND_1C,
        font_size: Pt = Pt(10),
        height: Inches = Inches(0.32),
        width: Optional[Inches] = None
    ) -> Any:
        """Отрисовывает компактный статус-бейдж в форме капсулы."""
        calc_width = width or Inches(len(text) * 0.1 + 0.35)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, calc_width, height
        )
        badge.adjustments[0] = 0.5  # Максимальное скругление в капсулу
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_color
        badge.line.color.rgb = text_color
        badge.line.width = Pt(0.75)

        tf = badge.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.04)
        tf.margin_bottom = Inches(0.04)
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_BODY
        p.font.size = font_size
        p.font.bold = True
        p.font.color.rgb = text_color
        return badge

    def add_card(
        self,
        slide: Any,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        title: Optional[str] = None,
        body: Optional[str] = None,
        bullets: Optional[List[str]] = None,
        badge_text: Optional[str] = None,
        badge_color: RGBColor = COLOR_BRAND_1C,
        border_color: RGBColor = COLOR_CARD_BORDER,
        fill_color: RGBColor = COLOR_CARD_BG,
        title_color: RGBColor = COLOR_TEXT_PRIMARY,
        body_font_size: Pt = Pt(11.5)
    ) -> Any:
        """Создает стильную модульную карточку с контуром и внутренними отступами."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.adjustments[0] = 0.04  # Аккуратное современное скругление 4%
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.0)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.24)
        tf.margin_right = Inches(0.24)
        tf.margin_top = Inches(0.22)
        tf.margin_bottom = Inches(0.22)

        has_title = False
        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = FONT_TITLE
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = title_color
            has_title = True

        if badge_text:
            self.add_badge(
                slide,
                left=left + width - Inches(len(badge_text) * 0.09 + 0.45),
                top=top + Inches(0.18),
                text=badge_text,
                bg_color=COLOR_CANVAS_BG,
                text_color=badge_color,
                font_size=Pt(9),
                height=Inches(0.26)
            )

        if body:
            p_body = tf.add_paragraph() if has_title else tf.paragraphs[0]
            p_body.text = body
            p_body.font.name = FONT_BODY
            p_body.font.size = body_font_size
            p_body.font.color.rgb = COLOR_TEXT_MUTED
            p_body.space_before = Pt(6)

        if bullets:
            for b in bullets:
                pb = tf.add_paragraph()
                pb.text = f"•  {b}"
                pb.font.name = FONT_BODY
                pb.font.size = body_font_size
                pb.font.color.rgb = COLOR_TEXT_MUTED
                pb.space_before = Pt(4)

        return card

    def add_metric_card(
        self,
        slide: Any,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        value: str,
        label: str,
        subtext: Optional[str] = None,
        accent_color: RGBColor = COLOR_BRAND_1C,
        badge_text: Optional[str] = None
    ) -> Any:
        """Создает карточку с акцентной большой числовой метрикой (KPI Card)."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.adjustments[0] = 0.04
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.0)

        if badge_text:
            self.add_badge(
                slide,
                left=left + width - Inches(len(badge_text) * 0.09 + 0.4),
                top=top + Inches(0.15),
                text=badge_text,
                bg_color=COLOR_CANVAS_BG,
                text_color=accent_color,
                font_size=Pt(8.5),
                height=Inches(0.24)
            )

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.24)
        tf.margin_right = Inches(0.24)
        tf.margin_top = Inches(0.22)
        tf.margin_bottom = Inches(0.2)

        # 1. Большая цифра
        pv = tf.paragraphs[0]
        pv.text = value
        pv.font.name = FONT_MONO
        pv.font.size = Pt(32)
        pv.font.bold = True
        pv.font.color.rgb = accent_color

        # 2. Подпись метрики
        pl = tf.add_paragraph()
        pl.text = label
        pl.font.name = FONT_TITLE
        pl.font.size = Pt(13)
        pl.font.bold = True
        pl.font.color.rgb = COLOR_TEXT_PRIMARY
        pl.space_before = Pt(4)

        # 3. Детальное пояснение
        if subtext:
            ps = tf.add_paragraph()
            ps.text = subtext
            ps.font.name = FONT_BODY
            ps.font.size = Pt(10.5)
            ps.font.color.rgb = COLOR_TEXT_MUTED
            ps.space_before = Pt(3)

        return card

    def add_pipeline(
        self,
        slide: Any,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        steps: List[Dict[str, Any]]
    ) -> None:
        """Отрисовывает горизонтальную архитектурную цепочку (Pipeline Flow)."""
        count = len(steps)
        if count == 0:
            return

        gap = Inches(0.22)
        arrow_width = Inches(0.25)
        total_gaps = (count - 1) * (gap + arrow_width)
        item_width = (width - total_gaps) / count

        cur_left = left
        for i, step in enumerate(steps):
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, cur_left, top, item_width, height
            )
            card.adjustments[0] = 0.05
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_CARD_BG
            card.line.color.rgb = COLOR_CARD_BORDER
            card.line.width = Pt(1.0)

            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.15)
            tf.margin_right = Inches(0.15)
            tf.margin_top = Inches(0.15)
            tf.margin_bottom = Inches(0.15)

            p_title = tf.paragraphs[0]
            p_title.text = step.get("title", f"Этап {i+1}")
            p_title.font.name = FONT_TITLE
            p_title.font.size = Pt(12)
            p_title.font.bold = True
            p_title.font.color.rgb = step.get("color", COLOR_BRAND_1C)

            if "desc" in step:
                p_desc = tf.add_paragraph()
                p_desc.text = step["desc"]
                p_desc.font.name = FONT_BODY
                p_desc.font.size = Pt(10)
                p_desc.font.color.rgb = COLOR_TEXT_MUTED
                p_desc.space_before = Pt(3)

            cur_left += item_width

            if i < count - 1:
                cur_left += gap / 2
                arrow_box = slide.shapes.add_textbox(
                    cur_left, top + height / 2 - Inches(0.25), arrow_width, Inches(0.5)
                )
                atf = arrow_box.text_frame
                atf.margin_left = 0
                atf.margin_right = 0
                ap = atf.paragraphs[0]
                ap.text = "→"
                ap.alignment = PP_ALIGN.CENTER
                ap.font.name = FONT_TITLE
                ap.font.size = Pt(20)
                ap.font.bold = True
                ap.font.color.rgb = COLOR_TECH_CYAN
                cur_left += arrow_width + gap / 2

    def add_table(
        self,
        slide: Any,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        headers: List[str],
        rows: List[List[str]],
        col_widths: Optional[List[Inches]] = None
    ) -> Any:
        """Отрисовывает контрастную корпоративную таблицу."""
        table_shape = slide.shapes.add_table(
            len(rows) + 1, len(headers), left, top, width, height
        )
        table = table_shape.table

        if col_widths and len(col_widths) == len(headers):
            for idx, w in enumerate(col_widths):
                table.columns[idx].width = w

        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = h
            p.font.name = FONT_TITLE
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_BRAND_1C

        for row_idx, row in enumerate(rows):
            is_alt = (row_idx % 2 == 1)
            row_bg = RGBColor(0x1E, 0x29, 0x3B) if not is_alt else RGBColor(0x16, 0x20, 0x32)
            for col_idx, val in enumerate(row):
                cell = table.cell(row_idx + 1, col_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_bg
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = val
                p.font.name = FONT_BODY
                p.font.size = Pt(10)
                p.font.color.rgb = COLOR_TEXT_PRIMARY if ("**" in val or "Да" in val) else COLOR_TEXT_MUTED

        return table_shape

    def add_speaker_notes(self, slide: Any, notes_text: str) -> None:
        """Добавляет текст речи спикера в нативные заметки PowerPoint."""
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    def save(self, filepath: str) -> None:
        """Сохраняет презентацию в указанный .pptx файл."""
        self.prs.save(filepath)
